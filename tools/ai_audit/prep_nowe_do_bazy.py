"""Przygotowanie wdrożenia pytań z Excela (Desktop\\Nowe do bazy): wiersze z odpowiedzią w kolumnie G i niepustą kolumną E (Odp. C).
Kolumna M ('Pewność i źródło') jest wykorzystywana do wyjaśnień. Etapy: odczyt -> odsiew duplikatów względem bazy i partii -> wyjaśnienia (GPT 6 Luna, zgodne z kluczem z G).
Baza tylko do odczytu. Wynik: review/Nowe_do_bazy/wdrozenie.json ({"items": [...], "duplikaty": [...], "pominiete": [...]}).
Użycie: python prep_nowe_do_bazy.py "<plik.xlsx>" """
import concurrent.futures as cf, difflib, json, math, re, sqlite3, sys, time
from pathlib import Path

import openpyxl
import requests

ROOT = Path(__file__).resolve().parents[2]
XLSX = Path(sys.argv[1])
OUT = ROOT / "review" / "Nowe_do_bazy"
OUT.mkdir(parents=True, exist_ok=True)
KEY = re.search(r"OPENROUTER_API_KEY\s*=\s*(\S+)", (ROOT / ".env").read_text(encoding="utf-8")).group(1)
H = {"Authorization": "Bearer " + KEY}
DB = str(Path.home() / "Documents" / "memonic_app.db")
ws = openpyxl.load_workbook(XLSX).active
head = [c.value for c in ws[1]]
col = {name: i for i, name in enumerate(head)}
items, skipped = [], []
for r in range(2, ws.max_row + 1):
    v = [c.value for c in ws[r]]
    q = v[col["Pytanie"]]
    if not q:
        continue
    g = str(v[col["ODP (Twoja)"]] or "").strip().upper()
    if g not in ("A", "B", "C", "D"):
        continue
    letter_cols = [k for k in ("Odp. A", "Odp. B", "Odp. C", "Odp. D", "Odp. E", "Odp. F") if k in col]
    opts = [re.sub(r"^[A-F]\)\s*", "", str(v[col[k]]).strip()) for k in letter_cols if v[col[k]] not in (None, "")]
    cval = v[col["Odp. C"]] if "Odp. C" in col else None
    if cval in (None, "") or len(opts) < 3 or ord(g) - 65 >= len(opts) or len(opts) > 4:
        skipped.append({"wiersz": r, "powod": "pusta kolumna Odp. C / za mało lub za dużo odpowiedzi / klucz poza zakresem", "pytanie": q})
        continue
    note_col = "Pewność i źródło" if "Pewność i źródło" in col else ("Uwagi" if "Uwagi" in col else None)
    ai_col = "ODP AI (3 z 5)" if "ODP AI (3 z 5)" in col else None
    items.append(dict(wiersz=r, nr=v[0], question=str(q).strip(), options=opts, key=g, ai=(v[col[ai_col]] if ai_col else None),
                      note=str(v[col[note_col]] or "").strip() if note_col else ""))
print("do wdrożenia (G + E niepuste):", len(items), "| pominięte:", len(skipped))


def norm(s):
    return re.sub(r"[^\w]+", " ", (s or "").lower()).strip()


con = sqlite3.connect(f"file:{Path(DB).as_posix()}?mode=ro", uri=True)
con.row_factory = sqlite3.Row
rows = con.execute("select * from questions").fetchall()
db_full = [(r["id"], norm(r["pytanie"]) + " " + " ".join(norm(r["odp_" + k]) for k in "abcd" if r["odp_" + k])) for r in rows]
db_stem = [(r["id"], norm(r["pytanie"])) for r in rows]


def fuzzy(q):
    k, s = norm(q["question"]) + " " + " ".join(norm(o) for o in q["options"]), norm(q["question"])
    bk = (0.0, None)
    for rid, t in db_full:
        if abs(len(t) - len(k)) > 0.5 * max(len(t), len(k), 1):
            continue
        sm = difflib.SequenceMatcher(None, k, t)
        if sm.real_quick_ratio() < max(bk[0], 0.75) or sm.quick_ratio() < max(bk[0], 0.75):
            continue
        x = sm.ratio()
        if x > bk[0]:
            bk = (x, rid)
    bs = (0.0, None)
    for rid, t in db_stem:
        if abs(len(t) - len(s)) > 0.4 * max(len(t), len(s), 1):
            continue
        x = difflib.SequenceMatcher(None, s, t).ratio()
        if x > bs[0]:
            bs = (x, rid)
    return bk, bs


keep, dups = [], []
seen = {}
for q in items:
    k = norm(q["question"]) + "|" + "|".join(norm(o) for o in q["options"])
    if k in seen:
        dups.append({**q, "powod": f"powtórzone w arkuszu (wiersz {seen[k]})"})
        continue
    seen[k] = q["wiersz"]
    (kr, kid), (sr, sid) = fuzzy(q)
    if kr >= 0.8 or sr >= 0.88:
        dups.append({**q, "powod": f"już w bazie (ID {kid if kr >= 0.8 else sid}, podobieństwo {max(kr, sr):.2f})", "w_bazie": kid if kr >= 0.8 else sid})
    else:
        keep.append(q)
print("po odsiewie rozmytym:", len(keep), "| duplikaty:", len(dups))


def chat(model, system, user, max_tokens=1500):
    for attempt in range(4):
        try:
            r = requests.post("https://openrouter.ai/api/v1/chat/completions", headers=H, timeout=120,
                              json={"model": model, "max_tokens": max_tokens, "reasoning": {"effort": "low"},
                                    "messages": [{"role": "system", "content": system}, {"role": "user", "content": user}]})
            if r.status_code == 200:
                return (r.json()["choices"][0]["message"]["content"] or "").strip()
            time.sleep(2 * (attempt + 1))
        except Exception:
            time.sleep(2 * (attempt + 1))
    return None


# ---- embeddingi + ocena podobieństwa semantycznego (pytania inaczej sformułowane)
def embed(texts):
    out = []
    for i in range(0, len(texts), 100):
        for attempt in range(4):
            r = requests.post("https://openrouter.ai/api/v1/embeddings", headers=H, timeout=120,
                              json={"model": "openai/text-embedding-3-small", "input": [t[:1800] for t in texts[i:i + 100]]})
            if r.status_code == 200:
                out += [d["embedding"] for d in r.json()["data"]]
                break
            time.sleep(3 * (attempt + 1))
        else:
            sys.exit("embeddings: błąd")
    return out


def unit(v):
    n = math.sqrt(sum(x * x for x in v)) or 1.0
    return [x / n for x in v]


def full(r):
    return r["pytanie"] + " | " + " | ".join(r["odp_" + k] for k in "abcd" if r["odp_" + k])


dbv = list(zip(rows, map(unit, embed([full(r) for r in rows]))))
ev = list(map(unit, embed([q["question"] + " | " + " | ".join(q["options"]) for q in keep])))
JSYS = ("You compare a Polish exam question with candidate questions from an existing database. Decide whether any candidate is the SAME exam question "
        "(same meaning, even if the wording or the order of options differs; the correct option must mean the same thing). Questions that only share a topic are NOT the same. "
        'Reply with JSON only: {"match": <candidate id or null>}')


def judge(i):
    sims = sorted(((sum(a * b for a, b in zip(ev[i], v)), r) for r, v in dbv), key=lambda x: -x[0])[:4]
    cand = []
    for _, r in sims:
        cand.append(f"id={r['id']}: {r['pytanie']}\n" + "\n".join(f"  {k.upper()}) {r['odp_' + k]}" for k in "abcd" if r["odp_" + k]))
    txt = chat("deepseek/deepseek-v4.1-flash", JSYS, "NEW QUESTION:\n" + keep[i]["question"] + "\n" + "\n".join(f"  {chr(97 + k)}) {o}" for k, o in enumerate(keep[i]["options"])) +
               "\n\nDATABASE CANDIDATES:\n" + "\n".join(cand))
    m = re.search(r"\{.*\}", txt or "", re.S)
    try:
        return json.loads(m.group(0)).get("match")
    except Exception:
        return None


with cf.ThreadPoolExecutor(3) as ex:
    matches = list(ex.map(judge, range(len(keep))))
final = []
for q, m in zip(keep, matches):
    if m:
        dups.append({**q, "powod": f"już w bazie (ID {m}, wg oceny AI)", "w_bazie": m})
    else:
        final.append(q)
print("po odsiewie semantycznym:", len(final), "| duplikaty razem:", len(dups))

# ---- wyjaśnienia zgodne z kluczem z G
ESYS = ("Jesteś ekspertem prawa morskiego i zarządzania statkiem. Napisz po polsku zwięzłe wyjaśnienie (2-4 zdania) do pytania testowego. "
        "Wyjaśnienie MA UZASADNIAĆ podaną poprawną odpowiedź (klucz od użytkownika jest ostateczny - nie kwestionuj go i nie sugeruj innej odpowiedzi). "
        "Wykorzystaj notatkę (podstawa prawna), jeśli jest; nie wymyślaj numerów artykułów, których nie ma w notatce ani których nie jesteś pewien. "
        "Wspomnij krótko, dlaczego pozostałe odpowiedzi nie są właściwe, jeśli da się to uczciwie uzasadnić. "
        "Zacznij dokładnie od: 'Prawidłowa odpowiedź to X.' (X = litera klucza). Bez nagłówków i list.")


def explain(q):
    opts = "\n".join(f"{chr(65 + k)}) {o}" for k, o in enumerate(q["options"]))
    user = f"Pytanie: {q['question']}\n{opts}\nPoprawna odpowiedź (klucz): {q['key']}\nNotatka: {q['note'] or '(brak)'}"
    t = chat("openai/gpt-6-luna", ESYS, user)
    if t and not t.startswith(f"Prawidłowa odpowiedź to {q['key']}."):
        t = f"Prawidłowa odpowiedź to {q['key']}. " + re.sub(r"^Prawidłowa odpowiedź to [A-C]\.\s*", "", t)
    return t


with cf.ThreadPoolExecutor(3) as ex:
    for q, t in zip(final, ex.map(explain, final)):
        q["explanation"] = t
miss = [q["wiersz"] for q in final if not q.get("explanation")]
print("wyjaśnienia gotowe:", len(final) - len(miss), "| brak:", miss)
json.dump({"items": final, "duplikaty": dups, "pominiete": skipped}, open(OUT / "wdrozenie.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
