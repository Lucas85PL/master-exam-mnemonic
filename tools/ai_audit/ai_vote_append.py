"""Dopisuje pytania z pliku MD (format: 'N. pytanie' + linie 'A. ...', 'B. ...', 'C. ...') do istniejącego Excela z głosowaniem AI
(kolumny: Pytanie, Odp. A/B/C, ODP AI (3 z 5), ODP (Twoja), AI1..AI5) i uzupełnia głosy 5 modeli:
AI1, AI2 = DeepSeek flash x2, AI3 = GPT 6.1 Sol, AI4, AI5 = GPT 6 Luna x2 (wszystko przez OpenRouter). ODP AI = odpowiedź wybrana przez >= 3 z 5.
Kopia Excela przed zmianą trafia do review/CCC_kopia/. Użycie: python ai_vote_append.py "<pytania.md>" "<arkusz.xlsx>" """
import concurrent.futures as cf, collections, json, re, shutil, sys, time
from copy import copy
from pathlib import Path

import openpyxl
import requests

ROOT = Path(__file__).resolve().parents[2]
MD, XLSX = Path(sys.argv[1]), Path(sys.argv[2])
KEY = re.search(r"OPENROUTER_API_KEY\s*=\s*(\S+)", (ROOT / ".env").read_text(encoding="utf-8")).group(1)
H = {"Authorization": "Bearer " + KEY}
VOTERS = [("AI1", "deepseek/deepseek-v4.1-flash", {"temperature": 0.7}), ("AI2", "deepseek/deepseek-v4.1-flash", {"temperature": 0.7}),
          ("AI3", "openai/gpt-6.1-sol", {}), ("AI4", "openai/gpt-6-luna", {}), ("AI5", "openai/gpt-6-luna", {})]
SYSTEM = ("You are an expert in maritime law and ship management (Polish Maritime Code, Polish maritime administration law, UNCLOS, SOLAS, STCW, MARPOL, ISM, "
          "ILO MLC and related conventions). Answer a Polish multiple-choice exam question: choose the SINGLE best answer. "
          'Reply with JSON only: {"answer": "<letter>"}')

# ---- parsowanie MD
text = MD.read_text(encoding="utf-8")
items = []
for block in re.split(r"\n\s*\n", text.strip()):
    lines = [l.strip() for l in block.strip().split("\n") if l.strip()]
    if not lines:
        continue
    m = re.match(r"^(\d+)\.\s*(.+)$", lines[0])
    if not m:
        continue
    opts = {}
    for l in lines[1:]:
        o = re.match(r"^([A-Da-d])[.)]\s*(.+)$", l)
        if o:
            opts[o.group(1).upper()] = o.group(2).strip()
    if len(opts) >= 2:
        items.append(dict(nr=int(m.group(1)), question=m.group(2).strip(), options=[opts[k] for k in sorted(opts)]))
print("pytań w MD:", len(items))


def letters(n):
    return [chr(65 + i) for i in range(n)]


def prompt(q):
    return q["question"] + "\n" + "\n".join(f"{L}) {o}" for L, o in zip(letters(len(q["options"])), q["options"]))


def ask(args):
    idx, v = args
    name, model, extra = VOTERS[v]
    q = items[idx]
    for attempt in range(4):
        try:
            body = {"model": model, "max_tokens": 2000, "reasoning": {"effort": "low"},
                    "messages": [{"role": "system", "content": SYSTEM}, {"role": "user", "content": prompt(q)}], **extra}
            r = requests.post("https://openrouter.ai/api/v1/chat/completions", headers=H, json=body, timeout=120)
            if r.status_code == 200:
                m = re.search(r'"answer"\s*:\s*"([A-Za-z])"', r.json()["choices"][0]["message"]["content"] or "")
                if m and m.group(1).upper() in letters(len(q["options"])):
                    return idx, name, m.group(1).upper()
            time.sleep(2 * (attempt + 1))
        except Exception:
            time.sleep(2 * (attempt + 1))
    return idx, name, None


votes = {}
for _ in range(2):  # drugie przejście ponawia brakujące głosy
    todo = [(i, v) for i in range(len(items)) for v in range(len(VOTERS)) if votes.get(f"{i}|{VOTERS[v][0]}") is None]
    with cf.ThreadPoolExecutor(4) as ex:
        for i, name, ans in ex.map(ask, todo):
            votes[f"{i}|{name}"] = ans

# ---- zapis do Excela (kopia przed zmianą)
bak = ROOT / "review" / "CCC_kopia"
bak.mkdir(parents=True, exist_ok=True)
shutil.copy2(XLSX, bak / (XLSX.stem + "_przed_dopisaniem.xlsx"))
wb = openpyxl.load_workbook(XLSX)
ws = wb.active
last = max(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value not in (None, ""))
tpl = last  # wiersz wzorcowy dla stylu
existing = {re.sub(r"\W+", " ", str(ws.cell(r, 1).value).lower()).strip() for r in range(2, last + 1) if ws.cell(r, 1).value}
dups = []
r = last
no_major = 0
for i, q in enumerate(items):
    r += 1
    v = [votes.get(f"{i}|{n}") for n, _, _ in VOTERS]
    top = collections.Counter(x for x in v if x).most_common(1)
    odp_ai = top[0][0] if top and top[0][1] >= 3 else ""
    no_major += 0 if odp_ai else 1
    if re.sub(r"\W+", " ", q["question"].lower()).strip() in existing:
        dups.append(q["nr"])
    opts = q["options"] + [""] * (3 - len(q["options"]))
    row = [q["question"]] + [f"{L}) {o}" if o else "" for L, o in zip("ABC", opts[:3])] + [odp_ai, ""] + [x or "?" for x in v]
    for c, val in enumerate(row, 1):
        cell = ws.cell(r, c, val)
        src = ws.cell(tpl, c)
        cell.alignment = copy(src.alignment) if src.has_style else cell.alignment
        cell.font = copy(src.font)
        cell.fill = copy(src.fill)
        cell.border = copy(src.border)
wb.save(XLSX)
print(f"dopisano wierszy: {len(items)} (od {last + 1} do {r}) | bez większości 3/5: {no_major} | jednomyślnie 5/5: "
      f"{sum(1 for i in range(len(items)) if len({votes.get(f'{i}|{n}') for n, _, _ in VOTERS}) == 1)} | już były w arkuszu (ten sam tekst): {dups}")
