"""Excel z nowymi pytaniami z 'Testy z Prawa Morskiego' (review/Testy_online/nowe_nieznane.json) + głosowanie 5 modeli AI.
Kolumny: Pytanie | Odp. A..C(D) | Odp. z arkusza (drukowana) | ODP (Twoja, pusta) | AI1..AI5 | ODP AI (3 z 5) | Uwagi.
AI1, AI2 = DeepSeek flash x2, AI3 = GPT 6.1 Sol, AI4, AI5 = GPT 6 Luna x2 (OpenRouter). Najpierw tańsze modele, na końcu Sol (limit klucza bywa na wyczerpaniu).
Użycie: python testy_online_excel.py [plik_wyjściowy.xlsx]"""
import concurrent.futures as cf, collections, json, re, sys, time
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
D = ROOT / "review" / "Testy_online"
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path.home() / "Desktop" / "Testy_online_nowe_pytania_glosowanie_AI.xlsx"
KEY = re.search(r"OPENROUTER_API_KEY\s*=\s*(\S+)", (ROOT / ".env").read_text(encoding="utf-8")).group(1)
H = {"Authorization": "Bearer " + KEY}
VOTERS = [("AI1", "deepseek/deepseek-v4.1-flash", {"temperature": 0.7}), ("AI2", "deepseek/deepseek-v4.1-flash", {"temperature": 0.7}),
          ("AI3", "openai/gpt-6.1-sol", {}), ("AI4", "openai/gpt-6-luna", {}), ("AI5", "openai/gpt-6-luna", {})]
ORDER = [0, 1, 3, 4, 2]  # Sol na końcu
items = json.load(open(D / "nowe_nieznane.json", encoding="utf-8"))
CACHE = D / "glosy_ai.json"
votes = json.load(open(CACHE, encoding="utf-8")) if CACHE.exists() else {}
SYSTEM = ("You are an expert in maritime law and ship management (Polish Maritime Code, Polish maritime administration law, UNCLOS, SOLAS, STCW, MARPOL, ISM, "
          "ILO MLC and related conventions). Answer a Polish multiple-choice exam question: choose the SINGLE best answer. "
          'Reply with JSON only: {"answer": "<letter>"}')


def letters(n):
    return [chr(65 + i) for i in range(n)]


def prompt(q):
    return q["question"] + "\n" + "\n".join(f"{L}) {o}" for L, o in zip(letters(len(q["options"])), q["options"]))


credit_out = {"stop": False}


def ask(args):
    idx, v = args
    name, model, extra = VOTERS[v]
    q = items[idx]
    if credit_out["stop"]:
        return idx, name, None
    for attempt in range(3):
        try:
            body = {"model": model, "max_tokens": 2000, "reasoning": {"effort": "low"},
                    "messages": [{"role": "system", "content": SYSTEM}, {"role": "user", "content": prompt(q)}], **extra}
            r = requests.post("https://openrouter.ai/api/v1/chat/completions", headers=H, json=body, timeout=120)
            if r.status_code == 200:
                m = re.search(r'"answer"\s*:\s*"([A-Za-z])"', r.json()["choices"][0]["message"]["content"] or "")
                if m and m.group(1).upper() in letters(len(q["options"])):
                    return idx, name, m.group(1).upper()
            elif r.status_code in (402, 403):
                credit_out["stop"] = True
                return idx, name, None
            time.sleep(2 * (attempt + 1))
        except Exception:
            time.sleep(2 * (attempt + 1))
    return idx, name, None


for v in ORDER:
    todo = [(i, v) for i in range(len(items)) if votes.get(f"{i}|{VOTERS[v][0]}") is None]
    with cf.ThreadPoolExecutor(4) as ex:
        for i, name, ans in ex.map(ask, todo):
            votes[f"{i}|{name}"] = ans
    json.dump(votes, open(CACHE, "w", encoding="utf-8"))
    print(VOTERS[v][0], "gotowe | brak odpowiedzi:", sum(1 for i in range(len(items)) if votes.get(f"{i}|{VOTERS[v][0]}") is None), "| limit klucza wyczerpany:", credit_out["stop"])

maxopt = max(len(q["options"]) for q in items)
wb = Workbook()
ws = wb.active
ws.title = "Testy online - nowe"
head = ["Pytanie"] + [f"Odp. {L}" for L in letters(maxopt)] + ["Odp. z arkusza (druk, niepewna)", "ODP (Twoja)", "AI1", "AI2", "AI3", "AI4", "AI5", "ODP AI (3 z 5)", "Uwagi"]
ws.append(head)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="305496")
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
no_major = 0
for i, q in enumerate(items):
    v = [votes.get(f"{i}|{n}") for n, _, _ in VOTERS]
    top = collections.Counter(x for x in v if x).most_common(1)
    odp_ai = top[0][0] if top and top[0][1] >= 3 else ""
    no_major += 0 if odp_ai else 1
    opts = q["options"] + [""] * (maxopt - len(q["options"]))
    uw = []
    if re.match(r"^[…\.]|^[a-ząćęłńóśźż]", q["question"]) or len(q["question"]) < 25:
        uw.append("treść pytania może być ucięta w PDF")
    if q.get("zrodlo"):
        uw.append("źródło w PDF: " + q["zrodlo"])
    ws.append([q["question"]] + [f"{L}) {o}" if o else "" for L, o in zip(letters(maxopt), opts)] + [q.get("key") or "", ""] + [x or "?" for x in v] + [odp_ai, "; ".join(uw)])
widths = [60] + [34] * maxopt + [16, 14, 7, 7, 7, 7, 7, 14, 36]
for k, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(k)].width = w
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "B2"
wb.save(OUT)
print("zapisano:", OUT, "| pytań:", len(items), "| bez większości 3/5:", no_major)
