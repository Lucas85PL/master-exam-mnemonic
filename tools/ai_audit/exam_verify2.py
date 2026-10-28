"""Druga runda weryfikacji: pytania MEDIUM i NISKA z exam_verify.py przechodza przez 3 sedziow: 2 x Claude Sonnet 5.5 + 1 x GPT 6.1 Sol (reasoning medium).
Wszyscy 3 sa 'slepi' (warianty przetasowane roznymi ziarnami, bez naszej odpowiedzi), dostaja fragmenty przepisow. Baza aplikacji nie jest ruszana.
Wynik: review/Weryfikacja_2026-10-12/wyniki_runda2.json / .xlsx"""
import hashlib
import json
import random
import re
import sys
import threading
import time
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import requests

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
OUT = ROOT / "review" / "Weryfikacja_2026-10-12"
CAP = float(sys.argv[1]) if len(sys.argv) > 1 else 20.0
JUDGES = [("sonnet_a", "anthropic/claude-sonnet-5.5", 101), ("sonnet_b", "anthropic/claude-sonnet-5.5", 202), ("sol", "openai/gpt-6.1-sol", 303)]
_lock = threading.Lock()
COST = defaultdict(float)
CACHE = OUT / "cache_runda2.jsonl"
_cache = {}
if CACHE.exists():
    for line in CACHE.read_text(encoding="utf-8").splitlines():
        r = json.loads(line)
        _cache[r["ck"]] = r["res"]
KEY = next(l.split("=", 1)[1].strip() for l in (ROOT / ".env").read_text(encoding="utf-8").splitlines() if l.startswith("OPENROUTER_API_KEY="))


def norm(s):
    return re.sub(r"[\W_]+", " ", (s or "").lower()).strip()


def call(model, system, user, tag):
    ck = hashlib.md5(f"{model}|{system}|{user}|{tag}".encode()).hexdigest()
    with _lock:
        if ck in _cache:
            return _cache[ck]
        if sum(COST.values()) >= CAP:
            return {"error": "cap"}
    body = {"model": model, "max_tokens": 10000, "reasoning": {"effort": "medium"}, "usage": {"include": True},
            "messages": [{"role": "system", "content": system}, {"role": "user", "content": user}]}
    err = None
    for i in range(3):
        try:
            r = requests.post("https://openrouter.ai/api/v1/chat/completions", headers={"Authorization": f"Bearer {KEY}"}, json=body, timeout=300)
            if r.status_code != 200:
                err = f"HTTP {r.status_code} {r.text[:120]}"
                time.sleep(3 * (i + 1))
                continue
            d = r.json()
            txt = d["choices"][0]["message"]["content"] or ""
            res = json.loads(re.search(r"\{.*\}", txt, re.S).group(0))
            with _lock:
                COST[model] += (d.get("usage") or {}).get("cost") or 0
                _cache[ck] = res
                with CACHE.open("a", encoding="utf-8") as f:
                    f.write(json.dumps({"ck": ck, "res": res}, ensure_ascii=False) + "\n")
            return res
        except Exception as e:  # noqa: BLE001
            err = repr(e)
            time.sleep(3 * (i + 1))
    return {"error": err}


# ---- fragmenty przepisow (jak w exam_verify.py)
SRC = {"Kodeks morski": ROOT / "review" / "Egzamin_2026-03-30" / "kodeks_morski.txt",
       "Ustawa o pracy na morzu": ROOT / "review" / "Egzamin_2026-03-30" / "src" / "pracy_na_morzu.txt",
       "Ustawa o PKBWM": ROOT / "review" / "Egzamin_2026-03-30" / "src" / "pkbwm.txt"}
arts = []
for name, p in SRC.items():
    if p.exists():
        for ptxt in re.split(r"(?=\nArt\. \d+[a-z]?\. )", p.read_text(encoding="utf-8")):
            m = re.match(r"\n(Art\. \d+[a-z]?)\. ", ptxt)
            if m and len(ptxt) > 80:
                arts.append((name, m.group(1), re.sub(r"\s+", " ", ptxt)[:1600]))
stop = set("który która które jest to na w z do od po dla oraz lub się nie jako przez czy że ich jego jej są być może gdy tylko".split())


def toks(s):
    return {w[:6] for w in norm(s).split() if len(w) > 3 and w not in stop}


art_t = [toks(a[2]) for a in arts]
df = defaultdict(int)
for ts in art_t:
    for w in ts:
        df[w] += 1


def excerpts(q):
    qt = toks(q["pytanie"] + " " + " ".join(q["opcje"].values()))
    sc = sorted(((sum(1.0 / (1 + df[w] ** 0.5) for w in qt & ts), n, a, t) for (n, a, t), ts in zip(arts, art_t)), key=lambda x: -x[0])
    return [x for x in sc[:3] if x[0] > 0.6]


SYS = ("Jesteś ekspertem prawa morskiego i bezpieczeństwa żeglugi (egzamin kapitański, poziom zarządzania, Polska). "
       "Odpowiadasz na podstawie obowiązującego stanu prawnego: Kodeks morski, ustawa o bezpieczeństwie morskim, ustawa o pracy na morzu, "
       "ustawa o PKBWM, ustawa o obszarach morskich, SOLAS, STCW, MLC 2006, MARPOL, COLREG, ISM, ISPS, LL 66, UNCLOS, dyrektywy UE (PSC 2009/16, 2002/59), "
       "Paris MoU, FAL, konwencje IMO. Pytania pochodzą z egzaminów z lat 2017-2024; jeśli przepis zmienił się od tamtej pory, oceń wg stanu obecnego i zaznacz to. "
       "Dostajesz fragmenty ustaw (mogą być nieistotne - wtedy je ignoruj). To pytanie egzaminacyjne: wskaż odpowiedź, którą uzna za prawidłową oficjalny klucz "
       "(najtrafniejszy wariant), nawet gdy sformułowanie jest niedoskonałe; flagę 'ambiguous' ustaw tylko wtedy, gdy co najmniej dwa warianty są równie dobre "
       "lub treść pytania jest sprzeczna z przepisem. Nie zgaduj: gdy nie masz pewności co do przepisu, obniż pewność. "
       'Zwróć WYŁĄCZNIE JSON: {"answer":"litera","confidence":0-100,"legal_basis":"konkretny przepis/artykuł",'
       '"flaw":"none|ambiguous|outdated|incomplete|multiple_correct","note":"max 1 zdanie"}')


def shuffled(q, seed):
    rnd = random.Random(seed)
    ks = list(q["opcje"])
    vals = [q["opcje"][k] for k in ks]
    rnd.shuffle(vals)
    new = dict(zip(ks, vals))
    back = {nk: next(k for k, v in q["opcje"].items() if v == nv) for nk, nv in new.items()}
    return new, back


def judge(rec, name, model, seed):
    q = rec
    new, back = shuffled(q, seed + len(q["pytanie"]))
    text = f"Pytanie: {q['pytanie']}\n" + "\n".join(f"  {k}) {v}" for k, v in new.items())
    ex = excerpts(q)
    if ex:
        text += "\n\nFragmenty przepisów:\n" + "\n".join(f"[{n} {a}] {t}" for _, n, a, t in ex)
    v = call(model, SYS, text, name)
    a = v.get("answer")
    v["answer_orig"] = back.get(a.strip().upper()[:1]) if isinstance(a, str) and a else None
    return v


data = json.loads((OUT / "wyniki.json").read_text(encoding="utf-8"))
todo = [d for d in data if d["status"] in ("MEDIUM", "NISKA")]
print("pytań do rundy 2:", len(todo), flush=True)


def work(d):
    return d["id"], {name: judge(d, name, model, seed) for name, model, seed in JUDGES}


res = {}
with ThreadPoolExecutor(12) as ex:
    for n, (i, r) in enumerate(ex.map(work, todo), 1):
        res[i] = r
        if n % 20 == 0:
            print(f"{n}/{len(todo)} koszt " + ", ".join(f"{k.split('/')[1]} ${v:.2f}" for k, v in COST.items()), flush=True)

final = []
for d in data:
    if d["id"] not in res:
        continue
    r = res[d["id"]]
    prop = d["odpowiedz_wstepna"]
    ans = {k: v.get("answer_orig") for k, v in r.items()}
    conf = {k: (v.get("confidence") or 0) for k, v in r.items()}
    flaws = {k: v.get("flaw") for k, v in r.items()}
    agree = sum(1 for a in ans.values() if a == prop)
    other = Counter(a for a in ans.values() if a and a != prop)
    err = any("error" in v for v in r.values())
    amb = sum(1 for f in flaws.values() if f == "ambiguous")
    bad = any(f in ("multiple_correct", "incomplete") for f in flaws.values())
    if err:
        st = "BLAD_API"
    elif agree == 3 and min(conf.values()) >= 80 and amb == 0 and not bad:
        st = "HIGH"
    elif agree == 3 and min(conf.values()) >= 70 and amb <= 1 and not bad:
        st = "MEDIUM"
    elif other and other.most_common(1)[0][1] >= 2:
        st = "ODRZUCONE_ROZBIEZNE"
    else:
        st = "NISKA"
    final.append({"id": d["id"], "status_runda1": d["status"], "status_runda2": st, "odpowiedz_wstepna": prop, "odpowiedzi_sedziow": ans,
                  "pewnosci": conf, "flagi": flaws, "przedmiot": d["przedmiot"], "pytanie": d["pytanie"], "opcje": d["opcje"],
                  "podstawy": {k: (v.get("legal_basis") or "")[:120] for k, v in r.items()},
                  "uwagi": {k: (v.get("note") or "")[:160] for k, v in r.items()}})
(OUT / "wyniki_runda2.json").write_text(json.dumps(final, ensure_ascii=False, indent=1), encoding="utf-8")
print(Counter((f["status_runda1"], f["status_runda2"]) for f in final))
print({k: round(v, 2) for k, v in COST.items()})

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
wb = Workbook()
ws = wb.active
ws.title = "Runda 2"
ws.append(["ID", "Runda 1", "Runda 2", "Przedmiot", "Pytanie", "A", "B", "C", "D", "Odp. wstępna", "Sonnet A", "Sonnet B", "Sol 6.1",
           "Pewności (SA/SB/Sol)", "Flagi", "Podstawy", "Uwagi"])
fill = {"HIGH": "C6EFCE", "MEDIUM": "FFEB9C", "NISKA": "FFC7CE", "ODRZUCONE_ROZBIEZNE": "F4B084"}
for f in final:
    o, a, c = f["opcje"], f["odpowiedzi_sedziow"], f["pewnosci"]
    ws.append([f["id"], f["status_runda1"], f["status_runda2"], f["przedmiot"], f["pytanie"], o.get("A", ""), o.get("B", ""), o.get("C", ""), o.get("D", ""),
               f["odpowiedz_wstepna"], a["sonnet_a"], a["sonnet_b"], a["sol"], f"{c['sonnet_a']}/{c['sonnet_b']}/{c['sol']}",
               "/".join(str(f['flagi'][k]) for k in ("sonnet_a", "sonnet_b", "sol")), " | ".join(f["podstawy"].values()), " | ".join(f["uwagi"].values())])
for i, w in enumerate([9, 10, 14, 18, 60, 28, 28, 28, 28, 8, 8, 8, 8, 14, 26, 60, 70]):
    ws.column_dimensions[chr(65 + i)].width = w
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
    if row[2].value in fill:
        row[2].fill = PatternFill("solid", fgColor=fill[row[2].value])
for c in ws[1]:
    c.font = Font(bold=True)
ws.freeze_panes = "F2"
ws.auto_filter.ref = ws.dimensions
wb.save(OUT / "wyniki_runda2.xlsx")
print("zapisano")
