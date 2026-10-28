"""Głosowanie 5 modeli AI nad odpowiedzią do pytań ze skanów (review/Nowe_21_10/nowe_nieznane.json) i zapis do Excela.
Głosujący: AI1, AI2 = DeepSeek flash (x2), AI3 = GPT 6.1 Sol, AI4, AI5 = GPT 6 Luna (x2); wszyscy przez OpenRouter (brak klucza DeepSeek API).
ODP AI = odpowiedź wybrana przez >=3 z 5 modeli; ODP = pusta kolumna na decyzję użytkownika. Baza nie jest używana ani zmieniana.
Użycie: python ai_vote_excel.py [plik_wyjściowy.xlsx]"""
import concurrent.futures as cf, collections, json, re, sys, time
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
D = ROOT / "review" / "Nowe_21_10"
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path.home() / "Desktop" / "Nowe_pytania_21_10_glosowanie_AI.xlsx"
KEY = re.search(r"OPENROUTER_API_KEY\s*=\s*(\S+)", (ROOT / ".env").read_text(encoding="utf-8")).group(1)
H = {"Authorization": "Bearer " + KEY}
VOTERS = [("AI1", "deepseek/deepseek-v4.1-flash", {"temperature": 0.7}),
          ("AI2", "deepseek/deepseek-v4.1-flash", {"temperature": 0.7}),
          ("AI3", "openai/gpt-6.1-sol", {}),
          ("AI4", "openai/gpt-6-luna", {}),
          ("AI5", "openai/gpt-6-luna", {})]
items = json.load(open(D / "nowe_nieznane.json", encoding="utf-8"))
CACHE = D / "glosy_ai.json"
votes = json.load(open(CACHE, encoding="utf-8")) if CACHE.exists() else {}
SYSTEM = ("You are an expert in maritime law and ship management (Polish Maritime Code, Polish maritime administration law, UNCLOS, SOLAS, STCW, MARPOL, ISM, "
          "ILO MLC and related conventions). Answer a Polish multiple-choice exam question: choose the SINGLE best answer. "
          'Reply with JSON only: {"answer": "<letter>"}')


def letters(n):
    return [chr(65 + i) for i in range(n)]


def prompt(q):
    return q["question"] + "\n" + "\n".join(f"{L}) {o}" for L, o in zip(letters(len(q["options"])), q["options"]))


usage = collections.Counter()


def ask(args):
    idx, voter = args
    name, model, extra = VOTERS[voter]
    q = items[idx]
    for attempt in range(4):
        try:
            body = {"model": model, "max_tokens": 2000, "reasoning": {"effort": "low"},
                    "messages": [{"role": "system", "content": SYSTEM}, {"role": "user", "content": prompt(q)}], **extra}
            r = requests.post("https://openrouter.ai/api/v1/chat/completions", headers=H, json=body, timeout=120)
            if r.status_code == 200:
                js = r.json()
                u = js.get("usage", {})
                usage[model + ":in"] += u.get("prompt_tokens", 0)
                usage[model + ":out"] += u.get("completion_tokens", 0)
                m = re.search(r'"answer"\s*:\s*"([A-Za-z])"', js["choices"][0]["message"]["content"] or "")
                if m and m.group(1).upper() in letters(len(q["options"])):
                    return idx, name, m.group(1).upper()
            time.sleep(2 * (attempt + 1))
        except Exception:
            time.sleep(2 * (attempt + 1))
    return idx, name, None


todo = [(i, v) for i in range(len(items)) for v in range(len(VOTERS)) if votes.get(f"{i}|{VOTERS[v][0]}") is None]
print("do zapytania:", len(todo), "z", len(items) * len(VOTERS))
with cf.ThreadPoolExecutor(4) as ex:
    for k, (i, name, ans) in enumerate(ex.map(ask, todo), 1):
        votes[f"{i}|{name}"] = ans
        if k % 40 == 0:
            json.dump(votes, open(CACHE, "w", encoding="utf-8"))
json.dump(votes, open(CACHE, "w", encoding="utf-8"))
cost = {"deepseek/deepseek-v4.1-flash": (0.049, 0.375), "openai/gpt-6.1-sol": (2.0, 10.0), "openai/gpt-6-luna": (0.1, 0.5)}
print("szacowany koszt tej sesji USD:", round(sum(usage[m + ":in"] * p[0] / 1e6 + usage[m + ":out"] * p[1] / 1e6 for m, p in cost.items()), 3))

maxopt = max(len(q["options"]) for q in items)
wb = Workbook()
ws = wb.active
ws.title = "Nowe pytania 21.10"
head = ["Pytanie"] + [f"Odp. {L}" for L in letters(maxopt)] + ["ODP AI (3 z 5)", "ODP (Twoja)", "AI1", "AI2", "AI3", "AI4", "AI5", "Drukowana odp. w skanie (niepewna)"]
ws.append(head)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="305496")
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
no_major = 0
for i, q in enumerate(items):
    v = [votes.get(f"{i}|{n}") for n, _, _ in VOTERS]
    cnt = collections.Counter(x for x in v if x)
    top = cnt.most_common(1)
    odp_ai = top[0][0] if top and top[0][1] >= 3 else ""
    if not odp_ai:
        no_major += 1
    opts = q["options"] + [""] * (maxopt - len(q["options"]))
    ws.append([q["question"]] + [f"{L}) {o}" if o else "" for L, o in zip(letters(maxopt), opts)] + [odp_ai, ""] + [x or "?" for x in v] +
              [(q.get("printed_key") or "").upper()])
widths = [60] + [34] * maxopt + [14, 14, 7, 7, 7, 7, 7, 18]
for k, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(k)].width = w
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "B2"
wb.save(OUT)
agree = sum(1 for i in range(len(items)) if len({votes.get(f"{i}|{n}") for n, _, _ in VOTERS}) == 1)
print("zapisano:", OUT, "| pytań:", len(items), "| bez większości 3/5:", no_major, "| jednomyślnie 5/5:", agree)
