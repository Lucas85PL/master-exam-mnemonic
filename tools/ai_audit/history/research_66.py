"""Research 66 pytan z Excela: modele z wyszukiwaniem w sieci wskazuja przepis (SOLAS/STCW/Kodeks morski...)."""
import argparse, json, re, sqlite3, threading
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
import requests
from openpyxl import load_workbook

HERE = Path(__file__).parent
MODELS = {"sonnet": "anthropic/claude-sonnet-5.5", "gpt": "openai/gpt-6.1-sol"}
SYSTEM = (
    "Jesteś ekspertem prawa morskiego. Dostajesz pytanie testowe z oficjalnego arkusza Urzędu Morskiego "
    "(egzamin na Kapitana żeglugi wielkiej) i odpowiedź z arkusza. Użyj wyszukiwarki, aby znaleźć "
    "PRZEPIS, który rozstrzyga pytanie (SOLAS, STCW, Kodeks morski, MARPOL, MLC, COLREG, LL, ISM/ISPS, "
    "ustawy i rozporządzenia polskie, publikacje IMO/IALA itd.). Oceń, czy odpowiedź z arkusza jest zgodna "
    "z przepisem. Zwróć WYŁĄCZNIE JSON: "
    '{"verdict":"POTWIERDZA|ZAPRZECZA|NIEJEDNOZNACZNE|BRAK_ZRODLA","best_answer":"A|B|C|D|NONE",'
    '"source":"akt + numer przepisu, np. SOLAS II-1/3-6","quote":"krótki cytat lub parafraza przepisu (max 250 znaków)",'
    '"url":"adres źródła","comment":"max 2 zdania po polsku"}. '
    "Jeśli nie znajdziesz przepisu, wybierz BRAK_ZRODLA - nie zgaduj i nie zmyślaj numerów przepisów."
)

def key():
    for l in (HERE / ".env").read_text(encoding="utf-8").splitlines():
        if l.startswith("OPENROUTER_API_KEY="): return l.split("=", 1)[1].strip()

def ask(k, model, q):
    opts = "\n".join(f"{x}: {q[x]}" for x in "ABCD" if q[x])
    body = {"model": model, "max_tokens": 9000, "reasoning": {"effort": "medium"},
            "usage": {"include": True}, "plugins": [{"id": "web", "max_results": 4}],
            "messages": [{"role": "system", "content": SYSTEM},
                         {"role": "user", "content": f"Pytanie:\n{q['pytanie']}\n\n{opts}\n\nOdpowiedź z arkusza: {q['ans']}"}]}
    for _ in range(3):
        try:
            r = requests.post("https://openrouter.ai/api/v1/chat/completions", headers={"Authorization": f"Bearer {k}"}, json=body, timeout=300)
            if r.status_code != 200: continue
            d = r.json(); t = d["choices"][0]["message"]["content"] or ""
            return json.loads(re.search(r"\{.*\}", t, re.S).group(0)), (d.get("usage") or {}).get("cost", 0)
        except Exception: pass
    return {"verdict": "BLAD_API"}, 0

if __name__ == "__main__":
    ap = argparse.ArgumentParser(); ap.add_argument("--limit", type=int); ap.add_argument("--out", default="research_66.jsonl"); a = ap.parse_args()
    ids = [r[0].value for r in load_workbook(HERE.parent / "Master_do_weryfikacji_v2.xlsx").active.iter_rows(min_row=2) if r[0].value]
    c = sqlite3.connect("C:/Users/Lukasz/Documents/memonic_app.db"); c.row_factory = sqlite3.Row
    out = Path(a.out); done = {json.loads(l)["id"] for l in out.read_text(encoding="utf-8").splitlines()} if out.exists() else set()
    todo = [i for i in ids if i not in done][: a.limit]; k = key(); lock = threading.Lock(); tot = [0]
    rows = {i: c.execute("select * from questions where id=?", (i,)).fetchone() for i in todo}
    def work(i):
        r = rows[i]
        q = {"pytanie": r["pytanie"], "A": r["odp_a"], "B": r["odp_b"], "C": r["odp_c"], "D": r["odp_d"], "ans": r["poprawna_odpowiedz"]}
        res = {n: ask(k, m, q) for n, m in MODELS.items()}
        rec = {"id": i, "db": q["ans"], "sonnet": res["sonnet"][0], "gpt": res["gpt"][0], "cost": res["sonnet"][1] + res["gpt"][1]}
        with lock:
            out.open("a", encoding="utf-8").write(json.dumps(rec, ensure_ascii=False) + "\n"); tot[0] += rec["cost"]
    with ThreadPoolExecutor(2) as ex: list(ex.map(work, todo))
    print("gotowe", len(todo), "koszt", round(tot[0], 2))
