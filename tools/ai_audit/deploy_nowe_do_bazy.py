"""Wdrożenie pytań z review/Nowe_do_bazy/wdrozenie.json (wynik prep_nowe_do_bazy.py) do kategorii Master w obu bazach, jako NOWE karty nauki,
z wyjaśnieniami zgodnymi z kluczem użytkownika (kolumna G). Potem zapis kopii Excela bez wierszy, które miały odpowiedź w G ('do opracowania').
Użycie: python deploy_nowe_do_bazy.py "<plik.xlsx>" ; aplikacja musi być zamknięta; kopie obu baz przed zapisem."""
import hashlib, json, shutil, sqlite3, subprocess, sys, time
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
XLSX = Path(sys.argv[1])
DATA = json.load(open(ROOT / "review" / "Nowe_do_bazy" / "wdrozenie.json", encoding="utf-8"))
DBS = {"app": Path.home() / "Documents" / "memonic_app.db", "seed": ROOT / "app" / "assets" / "db" / "memonic_app_seed.db"}
if "memonic_app.exe" in subprocess.run(["tasklist"], capture_output=True, text=True).stdout:
    sys.exit("Aplikacja jest uruchomiona - zamknij ją i powtórz.")
MARK = "manual:user_verified:2026-10-22"


def h(r):
    return hashlib.sha256("|".join([r[1], r[2] or "", r[3] or "", r[4] or "", r[5] or "", r[6], "v6-best-option"]).encode("utf-8")).hexdigest()


ts = time.strftime("%Y%m%d_%H%M%S")
items = DATA["items"]
for label, path in DBS.items():
    shutil.copy2(path, (ROOT / "data" if label == "seed" else path.parent) / (path.name + ".bak_nowedobazy_" + ts))
    db = sqlite3.connect(path)
    start = db.execute("SELECT MAX(id) FROM questions").fetchone()[0] + 1
    now = int(time.time() * 1000)
    for k, q in enumerate(items):
        qid = start + k
        o = q["options"] + [None] * (4 - len(q["options"]))
        notes = ("klucz wskazany przez użytkownika" + (f" | notatka: {q['note']}" if q["note"] else ""))[:500]
        db.execute("INSERT INTO questions (id,kategoria,tematyka,typ_pytania,pytanie,odp_a,odp_b,odp_c,odp_d,poprawna_odpowiedz,poprawna_odpowiedz_tekst,"
                   "zrodlo_pliku,review_status,review_notes,study_hidden,ai_verified) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                   (qid, "Kapitan żeglugi wielkiej", "Prawo morskie", "testowe", q["question"], o[0], o[1], o[2], o[3], q["key"], q["options"]["ABCD".index(q["key"])],
                    f"{XLSX.name} (Nowe do bazy), wiersz {q['wiersz']}", "ok", notes, 0, MARK))
        db.execute("INSERT INTO srs_cards (question_id,repetitions,interval_days,ease_factor,due_ms,lapses,state,step) VALUES (?,?,?,?,?,?,?,?)",
                   (qid, 0, 0, 2.5, now, 0, 0, 0))
        row = db.execute("SELECT id,pytanie,odp_a,odp_b,odp_c,odp_d,poprawna_odpowiedz FROM questions WHERE id=?", (qid,)).fetchone()
        db.execute("INSERT INTO explanation_cache VALUES (?,?,?,?,?)", (qid, h(row), q["explanation"], "gpt-6-luna (wyjaśnienie wg klucza użytkownika)", now))
    db.commit()
    print(label, "ID", start, "-", start + len(items) - 1, "| razem pytań", db.execute("SELECT COUNT(*) FROM questions").fetchone()[0])
    db.close()

# kopia Excela bez wierszy z odpowiedzią w kolumnie G
wb = openpyxl.load_workbook(XLSX)
ws = wb.active
head = [c.value for c in ws[1]]
g = head.index("ODP (Twoja)") + 1
removed = []
skipped_rows = {x["wiersz"] for x in DATA.get("pominiete", [])}  # G wypełnione, ale pusta kolumna E - zostają w pliku 'do opracowania'
for r in range(ws.max_row, 1, -1):
    if str(ws.cell(r, g).value or "").strip() and r not in skipped_rows:
        removed.append(r)
        ws.delete_rows(r)
out = XLSX.with_name(sys.argv[2]) if len(sys.argv) > 2 else XLSX.with_name(XLSX.stem + " - do opracowania.xlsx")
wb.save(out)
left = sum(1 for r in range(2, ws.max_row + 1) if ws.cell(r, 2).value)
print("kopia Excela:", out.name, "| usunięto wierszy z odpowiedzią w G:", len(removed), "| zostało 'do opracowania':", left)
